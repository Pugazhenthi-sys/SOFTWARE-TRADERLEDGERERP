import os
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from config import Config


def generate_bill_pdf(bill):
    path = os.path.join(Config.BILLS_DIR, f"{bill['bill_number']}.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Heading1"], alignment=TA_CENTER, textColor=colors.HexColor("#0F172A"))
    sub_style = ParagraphStyle("sub", parent=styles["Normal"], alignment=TA_CENTER, textColor=colors.HexColor("#475569"))

    biz = bill.get("business") or {}
    elements = []
    elements.append(Paragraph(biz.get("business_name") or "TradeLedger ERP", title_style))
    addr_line = ", ".join(filter(None, [biz.get("address"), biz.get("mobile_number"), biz.get("gst_number") and f"GSTIN: {biz.get('gst_number')}"]))
    if addr_line:
        elements.append(Paragraph(addr_line, sub_style))
    elements.append(Spacer(1, 10 * mm))

    meta_table = Table([
        ["Bill No:", bill["bill_number"], "Date:", bill["bill_date"]],
        ["Customer:", bill["customer_name"], "Vehicle No:", bill.get("vehicle_number") or "-"],
        ["Mobile:", bill.get("customer_mobile") or "-", "Address:", bill.get("customer_address") or "-"],
    ], colWidths=[25 * mm, 65 * mm, 25 * mm, 65 * mm])
    meta_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 8 * mm))

    data_rows = [["Product", "Tons", "Qty (kg)", "Bags", "Selling Rate", "Purchase Rate", "Duty/kg", "Sales Amount"]]
    data_rows.append([
        bill["product_name"], f"{bill['tons']:.3f}", f"{bill['quantity_kg']:.0f}", str(bill["bags"]),
        f"{bill['selling_rate']:.2f}", f"{bill['purchase_rate']:.2f}", f"{bill['duty_per_kg']:.2f}",
        f"{bill['sales_amount']:,.2f}",
    ])
    item_table = Table(data_rows, colWidths=[28 * mm, 15 * mm, 20 * mm, 14 * mm, 22 * mm, 22 * mm, 16 * mm, 26 * mm])
    item_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F766E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(item_table)
    elements.append(Spacer(1, 8 * mm))

    summary_style = ParagraphStyle("summary_right", parent=styles["Normal"], alignment=TA_RIGHT)
    summary_rows = [
        ["Purchase Amount", f"Rs. {bill['purchase_amount']:,.2f}"],
        ["Duty Amount", f"Rs. {bill['duty_amount']:,.2f}"],
        ["Gross Profit", f"Rs. {bill['gross_profit']:,.2f}"],
        ["Net Profit", f"Rs. {bill['net_profit']:,.2f}"],
        ["Previous Balance", f"Rs. {bill['previous_balance']:,.2f}"],
        ["Final Balance", f"Rs. {bill['final_balance']:,.2f}"],
    ]
    summary_table = Table(summary_rows, colWidths=[130 * mm, 33 * mm], hAlign="RIGHT")
    summary_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.HexColor("#0F172A")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(summary_table)

    if bill.get("remarks"):
        elements.append(Spacer(1, 6 * mm))
        elements.append(Paragraph(f"<b>Remarks:</b> {bill['remarks']}", styles["Normal"]))

    doc.build(elements)
    return path


def generate_bill_excel(bill):
    path = os.path.join(Config.BILLS_DIR, f"{bill['bill_number']}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Bill"

    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    bold = Font(bold=True)
    white_bold = Font(bold=True, color="FFFFFF")

    ws["A1"] = (bill.get("business") or {}).get("business_name") or "TradeLedger ERP"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")

    rows = [
        ("Bill No", bill["bill_number"], "Date", bill["bill_date"]),
        ("Customer", bill["customer_name"], "Vehicle No", bill.get("vehicle_number") or "-"),
        ("Mobile", bill.get("customer_mobile") or "-", "", ""),
    ]
    r = 3
    for a, b, c, d in rows:
        ws.cell(row=r, column=1, value=a).font = bold
        ws.cell(row=r, column=2, value=b)
        ws.cell(row=r, column=3, value=c).font = bold
        ws.cell(row=r, column=4, value=d)
        r += 1

    r += 1
    headers = ["Product", "Tons", "Qty (kg)", "Bags", "Selling Rate", "Purchase Rate", "Duty/kg", "Sales Amount"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = white_bold
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")
    r += 1
    ws.append([]) if False else None
    values = [bill["product_name"], bill["tons"], bill["quantity_kg"], bill["bags"],
              bill["selling_rate"], bill["purchase_rate"], bill["duty_per_kg"], bill["sales_amount"]]
    for i, v in enumerate(values, start=1):
        ws.cell(row=r, column=i, value=v)
    r += 2

    summary = [
        ("Purchase Amount", bill["purchase_amount"]),
        ("Duty Amount", bill["duty_amount"]),
        ("Gross Profit", bill["gross_profit"]),
        ("Net Profit", bill["net_profit"]),
        ("Previous Balance", bill["previous_balance"]),
        ("Final Balance", bill["final_balance"]),
    ]
    for label, val in summary:
        ws.cell(row=r, column=1, value=label).font = bold
        ws.cell(row=r, column=2, value=val)
        r += 1

    for col, width in zip("ABCDEFGH", [20, 16, 14, 10, 14, 14, 12, 16]):
        ws.column_dimensions[col].width = width

    wb.save(path)
    return path
