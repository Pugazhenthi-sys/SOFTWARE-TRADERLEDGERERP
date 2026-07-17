import os
import datetime
from flask import Blueprint, request, jsonify, send_file
from database import db_cursor
from auth_utils import jwt_required
from config import Config
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER

bp = Blueprint("reports", __name__, url_prefix="/api/reports")


def _resolve_range(period, start_date, end_date):
    today = datetime.date.today()
    if period == "daily":
        return today.isoformat(), today.isoformat()
    if period == "weekly":
        start = today - datetime.timedelta(days=today.weekday())
        return start.isoformat(), today.isoformat()
    if period == "monthly":
        return today.replace(day=1).isoformat(), today.isoformat()
    if period == "yearly":
        return today.replace(month=1, day=1).isoformat(), today.isoformat()
    return start_date or today.replace(day=1).isoformat(), end_date or today.isoformat()


def _sales_report_data(start_date, end_date):
    with db_cursor() as cur:
        cur.execute("""
            SELECT bills.bill_number, bills.bill_date, customers.name as customer_name,
                   products.name as product_name, bills.tons, bills.sales_amount,
                   bills.purchase_amount, bills.duty_amount, bills.gross_profit, bills.net_profit
            FROM bills
            JOIN customers ON customers.id = bills.customer_id
            JOIN products ON products.id = bills.product_id
            WHERE bills.bill_date BETWEEN ? AND ?
            ORDER BY bills.bill_date
        """, (start_date, end_date))
        bills = [dict(r) for r in cur.fetchall()]

        cur.execute("""
            SELECT COALESCE(SUM(sales_amount),0) sales, COALESCE(SUM(purchase_amount),0) purchase,
                   COALESCE(SUM(duty_amount),0) duty, COALESCE(SUM(gross_profit),0) gross_profit,
                   COALESCE(SUM(net_profit),0) net_profit, COUNT(*) bill_count
            FROM bills WHERE bill_date BETWEEN ? AND ?
        """, (start_date, end_date))
        totals = dict(cur.fetchone())

        cur.execute("""
            SELECT COALESCE(SUM(amount),0) total FROM expenses WHERE expense_date BETWEEN ? AND ?
        """, (start_date, end_date))
        totals["total_expenses"] = cur.fetchone()["total"]

    return bills, totals


@bp.route("/sales", methods=["GET"])
@jwt_required
def sales_report():
    period = request.args.get("period", "monthly")
    start_date, end_date = _resolve_range(period, request.args.get("start_date"), request.args.get("end_date"))
    bills, totals = _sales_report_data(start_date, end_date)
    return jsonify({"start_date": start_date, "end_date": end_date, "bills": bills, "totals": totals})


@bp.route("/sales/pdf", methods=["GET"])
@jwt_required
def sales_report_pdf():
    period = request.args.get("period", "monthly")
    start_date, end_date = _resolve_range(period, request.args.get("start_date"), request.args.get("end_date"))
    bills, totals = _sales_report_data(start_date, end_date)

    path = os.path.join(Config.REPORTS_DIR, f"sales_report_{start_date}_to_{end_date}.pdf")
    doc = SimpleDocTemplate(path, pagesize=landscape(A4), topMargin=14 * mm, bottomMargin=14 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Heading1"], alignment=TA_CENTER)
    elements = [Paragraph("TradeLedger ERP - Sales Report", title_style),
                Paragraph(f"{start_date} to {end_date}", styles["Normal"]), Spacer(1, 6 * mm)]

    rows = [["Bill No", "Date", "Customer", "Product", "Tons", "Sales", "Purchase", "Duty", "Gross Profit", "Net Profit"]]
    for b in bills:
        rows.append([b["bill_number"], b["bill_date"], b["customer_name"], b["product_name"],
                     f"{b['tons']:.2f}", f"{b['sales_amount']:,.2f}", f"{b['purchase_amount']:,.2f}",
                     f"{b['duty_amount']:,.2f}", f"{b['gross_profit']:,.2f}", f"{b['net_profit']:,.2f}"])
    rows.append(["", "", "", "TOTAL", "", f"{totals['sales']:,.2f}", f"{totals['purchase']:,.2f}",
                 f"{totals['duty']:,.2f}", f"{totals['gross_profit']:,.2f}", f"{totals['net_profit']:,.2f}"])

    table = Table(rows, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F766E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F1F5F9")]),
    ]))
    elements.append(table)
    doc.build(elements)
    return send_file(path, as_attachment=True, download_name=os.path.basename(path))


@bp.route("/sales/excel", methods=["GET"])
@jwt_required
def sales_report_excel():
    period = request.args.get("period", "monthly")
    start_date, end_date = _resolve_range(period, request.args.get("start_date"), request.args.get("end_date"))
    bills, totals = _sales_report_data(start_date, end_date)

    path = os.path.join(Config.REPORTS_DIR, f"sales_report_{start_date}_to_{end_date}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Report"
    headers = ["Bill No", "Date", "Customer", "Product", "Tons", "Sales", "Purchase", "Duty", "Gross Profit", "Net Profit"]
    fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
    for r, b in enumerate(bills, start=2):
        ws.append  # noop to keep structure explicit
        ws.cell(row=r, column=1, value=b["bill_number"])
        ws.cell(row=r, column=2, value=b["bill_date"])
        ws.cell(row=r, column=3, value=b["customer_name"])
        ws.cell(row=r, column=4, value=b["product_name"])
        ws.cell(row=r, column=5, value=b["tons"])
        ws.cell(row=r, column=6, value=b["sales_amount"])
        ws.cell(row=r, column=7, value=b["purchase_amount"])
        ws.cell(row=r, column=8, value=b["duty_amount"])
        ws.cell(row=r, column=9, value=b["gross_profit"])
        ws.cell(row=r, column=10, value=b["net_profit"])
    total_row = len(bills) + 2
    ws.cell(row=total_row, column=4, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=6, value=totals["sales"]).font = Font(bold=True)
    ws.cell(row=total_row, column=7, value=totals["purchase"]).font = Font(bold=True)
    ws.cell(row=total_row, column=8, value=totals["duty"]).font = Font(bold=True)
    ws.cell(row=total_row, column=9, value=totals["gross_profit"]).font = Font(bold=True)
    ws.cell(row=total_row, column=10, value=totals["net_profit"]).font = Font(bold=True)
    for col, width in zip("ABCDEFGHIJ", [14, 12, 20, 14, 10, 14, 14, 12, 14, 14]):
        ws.column_dimensions[col].width = width
    wb.save(path)
    return send_file(path, as_attachment=True, download_name=os.path.basename(path))
